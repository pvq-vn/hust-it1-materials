from pptx import Presentation
from pptx.util import Pt

# Khởi tạo presentation
prs = Presentation()

# Layouts
title_slide_layout = prs.slide_layouts[0]
bullet_slide_layout = prs.slide_layouts[1]
title_only_layout = prs.slide_layouts[5]

# Slide 1: Title
slide_1 = prs.slides.add_slide(title_slide_layout)
title_1 = slide_1.shapes.title
subtitle_1 = slide_1.placeholders[1]
title_1.text = "Research Ethics\nGóc khuất của Gian lận và Đạo văn"
subtitle_1.text = "Sinh viên thực hiện: Phạm Văn Quyết"

# Slide 2: Hook
slide_2 = prs.slides.add_slide(title_only_layout)
title_2 = slide_2.shapes.title
title_2.text = "Science is built on trust"
# Có thể chèn ảnh báo cáo bị "RETRACTED" vào slide này sau

# Slide 3: 3 Tội đồ (FFP)
slide_3 = prs.slides.add_slide(bullet_slide_layout)
title_3 = slide_3.shapes.title
body_shape_3 = slide_3.placeholders[1]
title_3.text = "3 Ranh giới đỏ trong nghiên cứu (FFP)"
tf_3 = body_shape_3.text_frame
tf_3.text = "1. Fabrication: Bịa đặt dữ liệu"
p1 = tf_3.add_paragraph()
p1.text = "2. Falsification: Xuyên tạc, xào nấu dữ liệu"
p2 = tf_3.add_paragraph()
p2.text = "3. Plagiarism: Đạo văn, ăn cắp chất xám"

# Slide 4: Case Study 1
slide_4 = prs.slides.add_slide(bullet_slide_layout)
title_4 = slide_4.shapes.title
body_shape_4 = slide_4.placeholders[1]
title_4.text = "Case Study: Cú lừa Tế bào gốc"
tf_4 = body_shape_4.text_frame
tf_4.text = "Nhân vật: Haruko Obokata (2014)"
p3 = tf_4.add_paragraph()
p3.text = "Sự kiện: Bịa đặt kết quả tạo tế bào gốc trên tạp chí Nature."
p4 = tf_4.add_paragraph()
p4.text = "Hậu quả: Bài báo bị rút, mất việc, đồng nghiệp tự sát."

# Slide 5: Case Study 2
slide_5 = prs.slides.add_slide(bullet_slide_layout)
title_5 = slide_5.shapes.title
body_shape_5 = slide_5.placeholders[1]
title_5.text = "Case Study: Vắc-xin & Bệnh Tự kỷ"
tf_5 = body_shape_5.text_frame
tf_5.text = "Nhân vật: Bác sĩ Andrew Wakefield (1998)"
p5 = tf_5.add_paragraph()
p5.text = "Sự kiện: Nhận tiền để làm giả dữ liệu vắc-xin gây tự kỷ."
p6 = tf_5.add_paragraph()
p6.text = "Hậu quả: Dịch sởi bùng phát do phụ huynh tẩy chay vắc-xin."

# Slide 6: Kỹ năng sinh tồn
slide_6 = prs.slides.add_slide(bullet_slide_layout)
title_6 = slide_6.shapes.title
body_shape_6 = slide_6.placeholders[1]
title_6.text = "Quote vs Paraphrase"
tf_6 = body_shape_6.text_frame
tf_6.text = "Quote (Trích dẫn trực tiếp): Dùng ngoặc kép, trích nguồn rõ ràng, giữ nguyên 100% từ ngữ."
p7 = tf_6.add_paragraph()
p7.text = "Paraphrase (Diễn đạt lại): Giữ nguyên ý, đổi cấu trúc câu, đổi từ vựng, trích nguồn đầy đủ."

# Slide 7: Take-away Message
slide_7 = prs.slides.add_slide(title_only_layout)
title_7 = slide_7.shapes.title
title_7.text = "Be honest, trustworthy, fair."

# Lưu file
prs.save('TWP_Research_Ethics.pptx')
print("Đã tạo xong file TWP_Research_Ethics.pptx!")